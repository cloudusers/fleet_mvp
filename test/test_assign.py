from __future__ import annotations

import unittest
from dataclasses import replace
from datetime import datetime

import tempfile
from pathlib import Path

from openpyxl import load_workbook

from fleet_mvp.assign import GroupALNS, format_dispatch, greedy_assign, rank_candidates, search_cost
from fleet_mvp.export_result import write_dispatch_xlsx
from fleet_mvp.execute import simulate_solution
from fleet_mvp.io_util import default_dataset
from fleet_mvp.loss import solution_cost


def _pair_empties(sol, kind_a: str, kind_b: str):
    sim = simulate_solution(sol)
    out = []
    for plan, ds in zip(sol.plans, sim.driver_sims):
        for prev, nxt, exe in zip(plan.groups, plan.groups[1:], ds.executions[1:]):
            if prev.kind == kind_a and nxt.kind == kind_b:
                out.append((prev.gid, nxt.gid, exe.empty_km))
    return out


class AssignTests(unittest.TestCase):
    def test_groups_are_not_split_or_duplicated(self):
        drivers, groups = default_dataset()
        sol = greedy_assign(drivers, groups)
        gids = sol.all_gids()
        self.assertEqual(sorted(gids), sorted(g.gid for g in groups))
        self.assertEqual(len(gids), len(set(gids)))
        self.assertTrue(simulate_solution(sol).feasible)

    def test_locked_current_group_stays_first(self):
        drivers, groups = default_dataset()
        sol = greedy_assign(drivers, groups)
        self.assertEqual(sol.locked_gids, ["G03"])
        d3 = next(p for p in sol.plans if p.driver.did == "D3")
        self.assertEqual(d3.groups[0].gid, "G03")
        sim = simulate_solution(sol)
        self.assertTrue(sim.feasible)
        d3s = next(ds for ds in sim.driver_sims if ds.did == "D3")
        self.assertEqual(d3s.executions[0].empty_km, 0.0)
        self.assertEqual(d3s.executions[0].reason, "in_progress")

    def test_alns_keeps_fixed_stop_order(self):
        drivers, groups = default_dataset()
        original = {g.gid: tuple(s.oid for s in g.stops) for g in groups}
        sol = GroupALNS(drivers, groups, max_iter=15, seed=2).run(verbose=False)
        for group in sol.assigned_groups() + sol.unassigned:
            self.assertEqual(tuple(s.oid for s in group.stops), original[group.gid])
        for gid in sol.locked_gids:
            self.assertTrue(any(g.gid == gid for g in sol.assigned_groups()))

    def test_normal_cover_or_explicit_unassigned(self):
        drivers, groups = default_dataset()
        sol = greedy_assign(drivers, groups)
        self.assertEqual(len(sol.assigned_groups()) + len(sol.unassigned), len(groups))
        self.assertLess(solution_cost(sol), 1e12)

    def test_alns_recovers_driver_stolen_by_greedy(self):
        drivers, groups = default_dataset(trap=True)
        greedy = greedy_assign(drivers, groups)
        self.assertEqual([g.gid for g in greedy.unassigned], ["T02"])
        d2 = next(p for p in greedy.plans if p.driver.did == "D2")
        self.assertEqual([g.gid for g in d2.groups], ["T01"])
        alns = GroupALNS(drivers, groups, max_iter=20, seed=42).run(verbose=False)
        self.assertEqual(alns.unassigned, [])
        self.assertLess(solution_cost(alns), solution_cost(greedy))
        d2 = next(p for p in alns.plans if p.driver.did == "D2")
        d5 = next(p for p in alns.plans if p.driver.did == "D5")
        self.assertEqual([g.gid for g in d2.groups], ["T02"])
        self.assertEqual([g.gid for g in d5.groups], ["T01"])

    def test_peak_has_shortage(self):
        drivers, groups = default_dataset(peak=True)
        start = datetime(2026, 9, 20, 14, 0)
        sol = greedy_assign(drivers, groups, now=start)
        self.assertGreater(len(sol.unassigned), 0)
        self.assertLess(len(sol.assigned_groups()), len(groups))

    def test_candidates_prefer_nearby_driver_for_namba_group(self):
        drivers, groups = default_dataset()
        g02 = next(g for g in groups if g.gid == "G02")
        sol = greedy_assign(drivers, groups)
        ranked = rank_candidates(sol, g02, top=5)
        feasible = [did for did, cost, _ in ranked if cost < 1e12]
        self.assertIn("D2", feasible)
        self.assertEqual(ranked[0][0], "D2")

    def test_usj_group_prefers_d4(self):
        drivers, groups = default_dataset()
        g04 = next(g for g in groups if g.gid == "G04")
        sol = greedy_assign(drivers, groups)
        ranked = rank_candidates(sol, g04, top=5)
        self.assertEqual(ranked[0][0], "D4")

    def test_one_model_covers_three_operating_patterns(self):
        drivers, groups = default_dataset()
        sol = greedy_assign(drivers, groups)
        self.assertTrue(simulate_solution(sol).feasible)
        songji_songji = _pair_empties(sol, "to_station", "to_station")
        songji_jieji = _pair_empties(sol, "to_station", "from_station")
        jieji_jieji = _pair_empties(sol, "from_station", "from_station")
        jieji_songji = _pair_empties(sol, "from_station", "to_station")
        self.assertTrue(songji_songji, msg="早晨连续送机会出现 T→酒店空驶")
        self.assertGreater(max(km for *_, km in songji_songji), 2.0)
        self.assertTrue(jieji_jieji, msg="晚上连续接机会出现 酒店→T 空驶")
        self.assertGreater(max(km for *_, km in jieji_jieji), 2.0)
        self.assertTrue(songji_jieji or jieji_songji, msg="白天轮换应出现空驶≈0 的衔接")
        mixed_empty = [km for *_, km in songji_jieji + jieji_songji]
        self.assertLess(min(mixed_empty), 0.2)

    def test_search_cost_matches_full_loss(self):
        drivers, groups = default_dataset()
        now = min(d.free_at for d in drivers)
        sol = greedy_assign(drivers, groups, now=now)
        self.assertAlmostEqual(search_cost(sol, now=now), solution_cost(sol, now=now), places=5)
        drivers, groups = default_dataset(peak=True)
        start = datetime(2026, 9, 20, 14, 0)
        peak = greedy_assign(drivers, groups, now=start)
        self.assertGreater(len(peak.unassigned), 0)
        self.assertAlmostEqual(search_cost(peak, now=start), solution_cost(peak, now=start), places=5)

    def test_destroy_weights_follow_scores(self):
        drivers, groups = default_dataset(trap=True)
        alns = GroupALNS(drivers, groups, max_iter=1, seed=1)
        alns._uses = [4, 0, 1]
        alns._scores = [20.0, 0.0, 0.0]
        alns._adapt_weights()
        self.assertGreater(alns.w_destroy[0], 1.0)
        self.assertGreater(alns.w_destroy[0], alns.w_destroy[2])
        self.assertEqual(alns.w_destroy[1], 1.0)
        self.assertEqual(alns._uses, [0, 0, 0])
        self.assertGreaterEqual(min(alns.w_destroy), 0.05)

    def test_dispatch_log_and_workbook_name_the_driver(self):
        drivers, groups = default_dataset(trap=True)
        sol = greedy_assign(drivers, groups)
        text = format_dispatch(sol)
        self.assertIn("佐藤", text)
        self.assertIn("车号", text)
        self.assertIn("T02", text)
        with tempfile.TemporaryDirectory() as folder:
            first = write_dispatch_xlsx(sol, directory=Path(folder))
            second = write_dispatch_xlsx(sol, directory=Path(folder))
            self.assertNotEqual(first.name, second.name)
            self.assertTrue(first.exists())
            self.assertTrue(second.exists())
            sheet = load_workbook(first).active
            header = [cell.value for cell in sheet[1]]
            self.assertEqual(header[0], "司机")
            self.assertIn("配车单号", header)
            self.assertIn("地址", header)
            names = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
            self.assertIn("佐藤", names)

    def test_jieji_then_same_hotel_songji_has_zero_empty(self):
        drivers, groups = default_dataset()
        sol = greedy_assign(drivers, groups)
        pairs = [(a, b, km) for a, b, km in _pair_empties(sol, "from_station", "to_station") if {a, b} == {"J01", "G11"}]
        self.assertTrue(pairs)
        self.assertLess(pairs[0][2], 0.2)

    def test_duplicate_group_or_driver_is_rejected(self):
        drivers, groups = default_dataset()
        with self.assertRaises(ValueError):
            greedy_assign(drivers, [groups[0], groups[0]])
        with self.assertRaises(ValueError):
            greedy_assign([drivers[0], drivers[0]], groups[:1])
        busy = dict(current_gid=groups[0].gid, status="enroute", dest_lat=34.66, dest_lon=135.50)
        first = replace(drivers[0], **busy)
        second = replace(drivers[1], **busy)
        with self.assertRaises(ValueError):
            greedy_assign([first, second], [groups[0]])
        idle = replace(drivers[0], current_gid=groups[0].gid, status="idle")
        with self.assertRaises(ValueError):
            greedy_assign([idle], [groups[0]])


if __name__ == "__main__":
    unittest.main()
