# 每次派单写一份新的 xlsx，不覆盖上一次。

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from fleet_mvp.assign import _clock
from fleet_mvp.execute import simulate_solution
from fleet_mvp.models import Group, Solution

HEADERS = (
    "司机",
    "车号",
    "配车单号",
    "送迎",
    "建议派单",
    "建议出发",
    "最晚出发",
    "站序",
    "客人",
    "订单号",
    "人数",
    "地址",
    "纬度",
    "经度",
    "预估时间",
    "时间含义",
)


def result_dir() -> Path:
    return Path(__file__).resolve().parent / "result"


def allocate_result_stem(directory: Optional[Path] = None) -> Path:
    # 运力表和对比文档共用这一段文件名，只是后缀不同。
    folder = Path(directory) if directory is not None else result_dir()
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    stem = folder / f"运力派单-{stamp}"
    if stem.with_suffix(".xlsx").exists() or stem.with_suffix(".md").exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        stem = folder / f"运力派单-{stamp}"
    return stem


def _action(group: Group) -> str:
    return "接客" if group.is_to_station() else "送达"


def _rows(sol: Solution, now: Optional[datetime]):
    sim = simulate_solution(sol, now=now, policy="jit")
    rows = []
    if not sim.feasible:
        raise ValueError(f"方案不可行，不写空表：{sim.reason}")
    exe_by_driver = {ds.did: ds for ds in sim.driver_sims}
    for plan in sol.plans:
        ds = exe_by_driver.get(plan.driver.did)
        if ds is None:
            continue
        for group, exe in zip(plan.groups, ds.executions):
            for index, stop in enumerate(group.stops, start=1):
                rows.append(
                    [
                        plan.driver.name,
                        plan.driver.plate,
                        group.gid,
                        group.kind_label(),
                        "" if group.gid in sol.locked_gids else _clock(exe.dispatch_at),
                        "" if group.gid in sol.locked_gids else _clock(exe.depart),
                        "" if group.gid in sol.locked_gids else _clock(exe.latest_depart),
                        index,
                        stop.guest or stop.name or stop.oid,
                        stop.oid,
                        stop.people,
                        stop.address,
                        stop.lat,
                        stop.lon,
                        _clock(stop.eta),
                        _action(group),
                    ]
                )
    for group in sol.unassigned:
        for index, stop in enumerate(group.stops, start=1):
            rows.append(
                [
                    "未派",
                    "",
                    group.gid,
                    group.kind_label(),
                    "",
                    "",
                    "",
                    index,
                    stop.guest or stop.name or stop.oid,
                    stop.oid,
                    stop.people,
                    stop.address,
                    stop.lat,
                    stop.lon,
                    _clock(stop.eta),
                    _action(group),
                ]
            )
    return rows


def write_dispatch_xlsx(
    sol: Solution,
    now: Optional[datetime] = None,
    directory: Optional[Path] = None,
    path: Optional[Path] = None,
) -> Path:
    if path is None:
        path = allocate_result_stem(directory).with_suffix(".xlsx")
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "运力派单"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in _rows(sol, now):
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (12, 10, 22, 8, 16, 16, 16, 8, 18, 22, 8, 48, 14, 14, 16, 10)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    book.save(path)
    return path
